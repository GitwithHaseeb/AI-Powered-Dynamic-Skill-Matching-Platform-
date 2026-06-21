"""
Convert CHATBOT_FULL_DOCUMENTATION.md into a formatted Word (.docx) file.

Run (from backend/):
  .\\venv312\\Scripts\\python.exe scripts/convert_chatbot_doc_to_docx.py

Input : <repo>/CHATBOT_FULL_DOCUMENTATION.md
Output: <repo>/CHATBOT_FULL_DOCUMENTATION.docx
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

_BACKEND = Path(__file__).resolve().parent.parent
_REPO = _BACKEND.parent
SRC = _REPO / "CHATBOT_FULL_DOCUMENTATION.md"
OUT = _REPO / "CHATBOT_FULL_DOCUMENTATION.docx"

_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_CODE_RE = re.compile(r"`([^`]+)`")


def _add_runs_with_markup(paragraph, text: str) -> None:
    """Render **bold** and `code` inline markup into a python-docx paragraph."""
    # Split on bold first; within each segment handle inline code.
    pos = 0
    for m in _BOLD_RE.finditer(text):
        _add_code_segments(paragraph, text[pos : m.start()], bold=False)
        _add_code_segments(paragraph, m.group(1), bold=True)
        pos = m.end()
    _add_code_segments(paragraph, text[pos:], bold=False)


def _add_code_segments(paragraph, text: str, bold: bool) -> None:
    if not text:
        return
    pos = 0
    for m in _CODE_RE.finditer(text):
        before = text[pos : m.start()]
        if before:
            r = paragraph.add_run(before)
            r.bold = bold
        cr = paragraph.add_run(m.group(1))
        cr.bold = bold
        cr.font.name = "Consolas"
        cr.font.color.rgb = RGBColor(0xB0, 0x30, 0x60)
        pos = m.end()
    rest = text[pos:]
    if rest:
        r = paragraph.add_run(rest)
        r.bold = bold


def _strip_md_table_row(line: str) -> list[str]:
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells


def _is_table_divider(line: str) -> bool:
    s = line.strip().strip("|").replace(" ", "")
    return bool(s) and set(s) <= {"-", ":", "|"}


def main() -> None:
    if not SRC.is_file():
        raise SystemExit(f"Source not found: {SRC}")

    lines = SRC.read_text(encoding="utf-8").splitlines()

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    i = 0
    n = len(lines)
    in_code = False
    code_buffer: list[str] = []

    while i < n:
        raw = lines[i]
        line = raw.rstrip("\n")

        # Fenced code blocks
        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_buffer = []
            else:
                in_code = False
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_buffer))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            i += 1
            continue
        if in_code:
            code_buffer.append(line)
            i += 1
            continue

        stripped = line.strip()

        # Horizontal rule -> spacer
        if stripped in ("---", "***", "___"):
            doc.add_paragraph("")
            i += 1
            continue

        # Headings
        if stripped.startswith("#"):
            m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
            if m:
                level = min(len(m.group(1)), 4)
                text = m.group(2).strip()
                if level == 1:
                    h = doc.add_heading("", level=0 if text.lower().startswith("chatbot —") else 1)
                    _add_runs_with_markup(h, text)
                    if text.lower().startswith("chatbot —"):
                        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    h = doc.add_heading("", level=level - 0)
                    _add_runs_with_markup(h, text)
                i += 1
                continue

        # Tables (markdown)
        if stripped.startswith("|") and i + 1 < n and _is_table_divider(lines[i + 1]):
            header = _strip_md_table_row(line)
            body_rows: list[list[str]] = []
            j = i + 2
            while j < n and lines[j].strip().startswith("|"):
                if _is_table_divider(lines[j]):
                    j += 1
                    continue
                body_rows.append(_strip_md_table_row(lines[j]))
                j += 1
            table = doc.add_table(rows=1, cols=len(header))
            try:
                table.style = "Light Grid Accent 1"
            except Exception:
                table.style = "Table Grid"
            hdr = table.rows[0].cells
            for c, htext in enumerate(header):
                hdr[c].paragraphs[0].text = ""
                run = hdr[c].paragraphs[0].add_run(htext)
                run.bold = True
            for row in body_rows:
                cells = table.add_row().cells
                for c in range(len(header)):
                    val = row[c] if c < len(row) else ""
                    cells[c].paragraphs[0].text = ""
                    _add_runs_with_markup(cells[c].paragraphs[0], val)
            i = j
            doc.add_paragraph("")
            continue

        # Bullet list
        if re.match(r"^\s*[-*]\s+", line):
            text = re.sub(r"^\s*[-*]\s+", "", line)
            p = doc.add_paragraph(style="List Bullet")
            _add_runs_with_markup(p, text)
            i += 1
            continue

        # Numbered list
        if re.match(r"^\s*\d+\.\s+", line):
            text = re.sub(r"^\s*\d+\.\s+", "", line)
            p = doc.add_paragraph(style="List Number")
            _add_runs_with_markup(p, text)
            i += 1
            continue

        # Blank line
        if not stripped:
            i += 1
            continue

        # Plain paragraph
        p = doc.add_paragraph()
        _add_runs_with_markup(p, stripped)
        i += 1

    doc.save(str(OUT))
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    main()
