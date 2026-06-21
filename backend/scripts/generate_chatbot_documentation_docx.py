"""
Build a professional Word report for the Skill Mapping Platform chatbot.

Reads `backend/docs/chatbot_400_queries.txt` if present (tab-separated LABEL<TAB>QUERY).
Regenerate that file with:
  python scripts/run_chatbot_battery.py --email admin@demo.local --queries-txt docs/chatbot_400_queries.txt --queries-only

Output: `backend/docs/Chatbot_Platform_Report.docx`
"""
from __future__ import annotations

import sys
from datetime import datetime, timezone
from pathlib import Path

_BACKEND = Path(__file__).resolve().parent.parent
if str(_BACKEND) not in sys.path:
    sys.path.insert(0, str(_BACKEND))

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Pt  # noqa: E402

DOCS = _BACKEND / "docs"
QUERIES_FILE = DOCS / "chatbot_400_queries.txt"
OUT_FILE = DOCS / "Chatbot_Platform_Report.docx"


def _load_queries() -> list[tuple[str, str]]:
    if not QUERIES_FILE.is_file():
        return []
    rows: list[tuple[str, str]] = []
    for line in QUERIES_FILE.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line:
            continue
        if "\t" in line:
            lab, q = line.split("\t", 1)
            rows.append((lab.strip(), q.strip()))
        else:
            rows.append((f"Q{len(rows) + 1:03d}", line))
    return rows


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    queries = _load_queries()
    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    t = doc.add_heading("Skill Mapping Platform — Chatbot Technical Report", level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Document generated: {now}")
    doc.add_paragraph("Group reference: F25CS093 (as used in repository documentation).")
    doc.add_paragraph("")

    doc.add_heading("1. Executive summary", level=1)
    doc.add_paragraph(
        "The platform exposes a bilingual assistant (English and Roman Urdu) that answers "
        "questions about projects, tasks, teams, skills, and activity using the same MongoDB "
        "collections as the rest of the application. A deterministic rule engine resolves the "
        "majority of intents; optional local Ollama integration can rephrase or complete "
        "answers when the deployment is configured in hybrid mode. This report describes the "
        "integration from end to end and lists the four-hundred-query evaluation battery used "
        "to exercise natural-language coverage before demonstration or examination."
    )

    doc.add_heading("2. Integration surface", level=1)
    doc.add_paragraph(
        "Clients call the authenticated HTTP endpoint POST /chatbot/query with a JSON body "
        "containing the user message. The FastAPI route in backend/app/routes/chatbot.py "
        "loads the current user from the JWT, fetches scoped projects, tasks, users, activity "
        "logs, and a performance snapshot from MongoDB, then dispatches the message through "
        "the same resolution order used by automated tests: strict testing phrases, direct "
        "owner-style answers, the rule engine (answer_with_context in backend/app/nlp/chat_engine.py), "
        "friendly fallbacks, and reply stabilization for first-person task lists."
    )

    doc.add_heading("3. Live data from MongoDB", level=1)
    doc.add_paragraph(
        "The chatbot does not rely on a static roster. On each request it reads live documents "
        "from the configured database (Motor async client; connection settings in backend/.env). "
        "Typical collections include users (with secrets removed), projects, tasks, and "
        "activity_logs. Administrators receive a wider project and task slice than managers "
        "or developers; non-admin callers still receive tasks assigned to or created by them "
        "so teammate questions remain consistent with in-app visibility rules."
    )

    doc.add_heading("4. Optional Ollama (local large language model)", level=1)
    doc.add_paragraph(
        "When CHATBOT_MODE is set to hybrid, the backend may call a locally hosted Ollama "
        "instance (defaults: OLLAMA_BASE_URL http://127.0.0.1:11434, OLLAMA_MODEL such as "
        "gemma3:27b). The model receives structured context derived from MongoDB and must not "
        "invent users or assignments. In rules-only mode, responses remain fully deterministic "
        "without contacting Ollama."
    )

    doc.add_heading("5. Four-hundred-query quality battery", level=1)
    doc.add_paragraph(
        "The script backend/scripts/run_chatbot_battery.py generates two hundred English "
        "and two hundred Roman Urdu template instantiations (four hundred total) using rotating "
        "placeholders drawn from the same visible projects, developers, skills, and task titles "
        "as the impersonated user. Running the script with --out produces a full question-and-answer "
        "transcript; running with --queries-txt (and optionally --queries-only) exports the exact "
        "query strings for manual re-testing in the web UI or on localhost."
    )
    if not queries:
        doc.add_paragraph(
            "Appendix A is empty because chatbot_400_queries.txt was not found. Generate it with:\n"
            "cd backend && python scripts/run_chatbot_battery.py --email admin@demo.local "
            "--queries-txt docs/chatbot_400_queries.txt --queries-only"
        )
    else:
        doc.add_paragraph(f"Appendix A contains all {len(queries)} exported queries.")

    doc.add_heading("6. Reproducibility", level=1)
    doc.add_paragraph(
        "Prerequisites: Python environment with backend/requirements.txt installed, valid "
        "MONGODB_URL in backend/.env, and seeded or production-like data. Example commands:\n"
        "• Export query list only:\n"
        "  python scripts/run_chatbot_battery.py --email admin@demo.local "
        "--queries-txt docs/chatbot_400_queries.txt --queries-only\n"
        "• Full battery with report file:\n"
        "  python scripts/run_chatbot_battery.py --email admin@demo.local "
        "--queries-txt docs/chatbot_400_queries.txt --out docs/chatbot_400_battery_report.txt"
    )

    doc.add_heading("Appendix A — Query list (label and text)", level=1)
    if not queries:
        doc.add_paragraph("(No queries file — see section 5.)")
    else:
        for lab, q in queries:
            p = doc.add_paragraph(style="List Number")
            p.add_run(f"{lab}: ").bold = True
            p.add_run(q)

    doc.save(OUT_FILE)
    print(f"Wrote {OUT_FILE}")


if __name__ == "__main__":
    main()
