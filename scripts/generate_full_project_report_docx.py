"""
Generate a complete A-to-Z Word report for the Skill Mapping Platform.
Output: docs/Skill_Mapping_Platform_Complete_Report.docx
"""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
OUT_FILE = DOCS / "Skill_Mapping_Platform_Complete_Report.docx"


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def build_document() -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    now = datetime.now(timezone.utc).strftime("%d %B %Y")

    # ── Title Page ──
    title = doc.add_heading("AI-Powered Dynamic Skill Matching Platform", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("Complete Project Report (A to Z)")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].bold = True
    sub.runs[0].font.size = Pt(14)
    meta = doc.add_paragraph(f"Group: F25CS093 | BSCS Final Year Project\nGenerated: {now}")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    # ── Table of Contents placeholder ──
    add_heading(doc, "Table of Contents", 1)
    toc_items = [
        "1. Executive Summary",
        "2. Introduction and Project Background",
        "3. Problem Statement and Objectives",
        "4. Technology Stack",
        "5. System Architecture",
        "6. Repository Structure",
        "7. Database Design",
        "8. User Roles and Access Control",
        "9. Authentication and Security",
        "10. Backend Application (FastAPI)",
        "11. Machine Learning and NLP Components",
        "12. Frontend Application (React + Vite)",
        "13. Core Workflows",
        "14. AI Team Recommendations",
        "15. Task Management System",
        "16. Analytics and Reporting",
        "17. Hybrid Chatbot",
        "18. API Endpoints Overview",
        "19. Environment Configuration",
        "20. Installation and Running the Platform",
        "21. Demo Scenarios",
        "22. Security and Privacy Considerations",
        "23. Limitations and Future Enhancements",
        "24. Conclusion",
    ]
    add_numbered(doc, toc_items)
    doc.add_page_break()

    # ── 1. Executive Summary ──
    add_heading(doc, "1. Executive Summary", 1)
    add_para(
        doc,
        "The Skill Mapping Platform is an AI-powered web application designed for intelligent "
        "developer–project matching in software development teams. Built as a BSCS Final Year Project "
        "(Group F25CS093), the platform follows an SDS (Software Design Specification) compliant "
        "architecture and delivers a complete loop: document-based skill extraction, AI-driven team "
        "recommendations, manager approval, task assignment, adaptive skill profiling, analytics "
        "reporting, and a bilingual (English / Roman Urdu) chatbot grounded in live MongoDB data.",
    )
    add_para(
        doc,
        "The system consists of a FastAPI backend with MongoDB persistence, a React + Vite frontend, "
        "scikit-learn based skill matching, spaCy/document parsing for SRS files, optional local Ollama "
        "LLM integration for the chatbot, and role-based dashboards for Admin, Project Manager (PM), "
        "and Developer users.",
    )

    # ── 2. Introduction ──
    add_heading(doc, "2. Introduction and Project Background", 1)
    add_para(
        doc,
        "Modern software organizations face a recurring challenge: assigning the right developers to "
        "the right projects based on skills, availability, and experience. Manual team formation is "
        "slow, subjective, and often leads to skill gaps or overloaded developers. This platform "
        "automates that process using machine learning, natural language processing, and structured "
        "workflow management.",
    )
    add_heading(doc, "2.1 What the Platform Does", 2)
    add_bullets(
        doc,
        [
            "Parses Software Requirements Specification (SRS) documents to detect required skills.",
            "Matches developers to projects using TF-IDF cosine similarity and weighted scoring.",
            "Recommends stratified top-5 developer teams per project with human-readable explanations.",
            "Allows managers to approve AI recommendations and automatically assign starter tasks.",
            "Tracks tasks through a full lifecycle: assigned → in progress → submitted → completed.",
            "Updates developer skill proficiency when tasks are marked complete by a PM.",
            "Provides analytics dashboards with utilization metrics, skill gaps, and PDF export.",
            "Offers a hybrid chatbot that answers questions in English or Roman Urdu using live data.",
        ],
    )

    # ── 3. Problem Statement ──
    add_heading(doc, "3. Problem Statement and Objectives", 1)
    add_heading(doc, "3.1 Problem Statement", 2)
    add_para(
        doc,
        "Organizations lack an integrated system that connects document analysis, skill profiling, "
        "team recommendation, task tracking, and adaptive learning in one platform. Existing tools "
        "often handle only one aspect (e.g., project management OR skill tracking) without AI-driven "
        "matching or real-time conversational access to project data.",
    )
    add_heading(doc, "3.2 Project Objectives", 2)
    add_bullets(
        doc,
        [
            "Build an SDS-compliant full-stack application with JWT authentication and RBAC.",
            "Extract skills automatically from SRS/resume documents using NLP.",
            "Implement ML-based developer–project matching with explainable scores.",
            "Enable manager-controlled approval of AI team recommendations.",
            "Support fair task distribution based on developer workload.",
            "Provide adaptive skill updates when tasks are completed.",
            "Deliver analytics for team performance and skill gap identification.",
            "Integrate a privacy-preserving hybrid chatbot with local LLM support (Ollama).",
        ],
    )

    # ── 4. Technology Stack ──
    add_heading(doc, "4. Technology Stack", 1)
    add_heading(doc, "4.1 Backend", 2)
    add_bullets(
        doc,
        [
            "Python 3.11+ with FastAPI (async REST API framework)",
            "MongoDB with Motor (async driver) for data persistence",
            "Pydantic / pydantic-settings for configuration and validation",
            "JWT (python-jose) + bcrypt (passlib) for authentication",
            "scikit-learn (TF-IDF, cosine similarity) for skill matching",
            "spaCy, pdfplumber, python-docx, pytesseract for document parsing",
            "Optional TensorFlow ranker (tensorflow_rank.py) for advanced ranking",
            "Ollama integration for local LLM chatbot fallback",
            "ReportLab for PDF analytics report generation",
        ],
    )
    add_heading(doc, "4.2 Frontend", 2)
    add_bullets(
        doc,
        [
            "React 18 with React Router for SPA navigation",
            "Vite as the development server and build tool",
            "Tailwind CSS for styling",
            "Axios-based API service layer with JWT token handling",
            "Context API (AuthContext, ThemeContext) for global state",
            "Chart components for analytics visualization",
        ],
    )
    add_heading(doc, "4.3 Infrastructure and Tools", 2)
    add_bullets(
        doc,
        [
            "MongoDB (local, Docker, or MongoDB Atlas)",
            "Docker Compose for local MongoDB container",
            "Swagger UI / ReDoc for API documentation at /docs",
            "Ollama (optional) for local LLM inference",
            "VS Code launch configurations for full-stack debugging",
        ],
    )

    # ── 5. Architecture ──
    add_heading(doc, "5. System Architecture", 1)
    add_para(
        doc,
        "The platform follows a classic three-tier architecture: Presentation (React frontend), "
        "Application (FastAPI backend with ML/NLP services), and Data (MongoDB). The frontend "
        "communicates with the backend via REST API calls proxied through Vite during development.",
    )
    add_heading(doc, "5.1 Request Flow", 2)
    add_numbered(
        doc,
        [
            "User interacts with the React UI in the browser (typically http://localhost:5173).",
            "Frontend sends HTTP requests to /api/* which Vite proxies to FastAPI on port 8000.",
            "StripApiPrefixMiddleware strips the /api prefix so routes match backend paths.",
            "JWT middleware (get_current_user dependency) validates the Bearer token.",
            "Route handlers query or update MongoDB collections via async Motor client.",
            "ML/NLP services (SkillMatcher, DocumentParser, ChatEngine) process business logic.",
            "JSON responses are returned to the frontend for rendering.",
        ],
    )
    add_heading(doc, "5.2 Key Architectural Decisions", 2)
    add_bullets(
        doc,
        [
            "MongoDB as a flexible document store for users, projects, tasks, and recommendations.",
            "Async FastAPI for non-blocking I/O with MongoDB.",
            "Rule-first chatbot with optional LLM fallback for speed and data grounding.",
            "Demo project titles (Demo —*) hidden from production UI lists via demo_projects.py.",
            "Stratified top-5 team recommendations configurable via MAX_PROJECT_TEAM_SIZE.",
        ],
    )

    # ── 6. Repository Structure ──
    add_heading(doc, "6. Repository Structure", 1)
    add_para(doc, "The project is organized into the following main directories:")
    structure = [
        ("backend/", "FastAPI application, ML/NLP modules, seed scripts, requirements.txt"),
        ("backend/app/main.py", "Application entry point; registers routers and middleware"),
        ("backend/app/routes/", "API route modules (auth, users, projects, tasks, ml, chatbot, etc.)"),
        ("backend/app/models/", "Pydantic schemas for users, projects, and tasks"),
        ("backend/app/core/", "Security, document parser, skill matcher, match explanation"),
        ("backend/app/ml/", "scikit-learn vectors, skill extractor, TensorFlow ranker"),
        ("backend/app/nlp/", "Chat engine with rule-based and LLM paths"),
        ("backend/app/services/", "Activity log, analytics PDF, team roster, task descriptions"),
        ("backend/app/utils/", "MongoDB helpers, demo project filter, validators"),
        ("backend/docs/", "Chatbot battery files and generated reports"),
        ("backend/scripts/", "Utility scripts (chatbot battery, backfill, doc generation)"),
        ("frontend/", "React + Vite frontend application"),
        ("frontend/src/Components/", "UI components (Dashboards, TaskList, ChatBot, etc.)"),
        ("frontend/src/context/", "AuthContext and ThemeContext providers"),
        ("frontend/src/services/", "API client (api.js) with axios and token management"),
        ("docs/", "Project documentation and generated Word reports"),
        ("DEMO_SCRIPT.md", "Defense and viva demonstration checklist"),
    ]
    for path, desc in structure:
        p = doc.add_paragraph()
        run = p.add_run(f"{path} — ")
        run.bold = True
        p.add_run(desc)

    # ── 7. Database ──
    add_heading(doc, "7. Database Design", 1)
    add_para(
        doc,
        "MongoDB database name defaults to skill_mapping (configurable via MONGODB_DB_NAME). "
        "The following collections store the platform data:",
    )
    add_heading(doc, "7.1 Collections", 2)
    collections = [
        ("users", "User accounts with email, hashed password, role, skills array, availability, CGPA, experience"),
        ("projects", "Project metadata: title, description, require_skills, deadline, status, assigned_team, final_team, progress"),
        ("tasks", "Task records: title, description, status, assigned_to, project_id, skills_used, submission files"),
        ("recommendations", "Pending/approved AI team recommendation batches with candidate scores"),
        ("skills", "Optional skill taxonomy/reference data"),
        ("srs_documents", "Uploaded SRS document metadata and parsed content"),
        ("activity_logs", "Audit trail for user actions (approvals, status changes, etc.)"),
    ]
    for name, desc in collections:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(f"{name}: ")
        run.bold = True
        p.add_run(desc)
    add_heading(doc, "7.2 Connection Management", 2)
    add_para(
        doc,
        "database.py manages the AsyncIOMotorClient lifecycle. On application startup, db.connect() "
        "establishes the connection. On shutdown, db.disconnect() closes it. Helper functions like "
        "get_users_collection(), get_projects_collection(), and get_tasks_collection() provide "
        "typed access to each collection.",
    )

    # ── 8. Roles ──
    add_heading(doc, "8. User Roles and Access Control", 1)
    add_para(doc, "The platform implements Role-Based Access Control (RBAC) with three primary roles:")
    add_heading(doc, "8.1 Admin", 2)
    add_bullets(
        doc,
        [
            "Full organization-wide visibility: all projects, all tasks, all users.",
            "Access to analytics dashboard and developer directory.",
            "Can approve AI recommendations and manage any project.",
            "Widest chatbot data scope for demos and oversight.",
        ],
    )
    add_heading(doc, "8.2 Manager (Project Manager / PM)", 2)
    add_bullets(
        doc,
        [
            "Creates and manages own projects.",
            "Views AI developer recommendations for their projects.",
            "Approves AI teams and assigns tasks to developers.",
            "Reviews developer task submissions and marks tasks as completed.",
            "Accesses analytics for owned projects.",
            "Main project grid shows active projects; completed projects appear in a separate section.",
        ],
    )
    add_heading(doc, "8.3 Developer", 2)
    add_bullets(
        doc,
        [
            "Views assigned projects and tasks.",
            "Updates task status (in progress, submit for review).",
            "Cannot directly mark tasks as completed (PM approval required).",
            "Maintains skill profile updated automatically on task completion.",
            "Home dashboard shows completed projects summary; Tasks page focuses on workflow.",
            "Restricted from accessing other PMs' projects or developer listing endpoints.",
        ],
    )

    # ── 9. Auth ──
    add_heading(doc, "9. Authentication and Security", 1)
    add_heading(doc, "9.1 Registration and Login", 2)
    add_para(
        doc,
        "Users register via POST /auth/register with email, password, username, and optional role. "
        "Email validation uses Pydantic EmailStr. Login via POST /auth/login performs case-insensitive "
        "email lookup and bcrypt password verification. On success, a JWT access token is returned "
        "along with the user profile.",
    )
    add_heading(doc, "9.2 JWT Token Structure", 2)
    add_bullets(
        doc,
        [
            "Algorithm: HS256 with SECRET_KEY from environment.",
            "Claims: sub (email), user_id, role, exp (expiration).",
            "Default expiration: 7 days (ACCESS_TOKEN_EXPIRE_MINUTES=10080).",
            "Frontend stores token and sends Authorization: Bearer <token> on each request.",
        ],
    )
    add_heading(doc, "9.3 Protected Routes", 2)
    add_para(
        doc,
        "The get_current_user dependency in auth.py decodes the JWT, loads the user from MongoDB, "
        "and injects UserResponse into route handlers. Unauthorized or invalid tokens return 401. "
        "Role-specific restrictions return 403 when a user attempts an action outside their scope.",
    )

    # ── 10. Backend ──
    add_heading(doc, "10. Backend Application (FastAPI)", 1)
    add_para(
        doc,
        "The backend is organized into modular routers registered in main.py. Each router handles "
        "a specific domain of the application.",
    )
    routers = [
        ("auth.py", "/auth", "Register, login, GET /auth/me for current user profile"),
        ("users.py", "/users", "Developer listing, profile updates, skill management, resume upload"),
        ("projects.py", "/projects", "CRUD for projects, SRS upload, AI recommendations endpoint"),
        ("projects.team_details_router", "/projects", "Team details modal data (registered before generic project routes)"),
        ("tasks.py", "/tasks", "Task CRUD, submission workflow, status transitions, adaptive skills"),
        ("teams.py", "/teams", "Team listing and management aligned with project assigned_team"),
        ("analytics.py", "/analytics", "Dashboard metrics, utilization, skill gaps, PDF export"),
        ("chatbot.py", "/chatbot", "POST /chatbot/query — hybrid NLP chatbot"),
        ("recommendation_flow.py", "/recommendations", "Create, approve, reject AI team recommendations"),
        ("ml.py", "/ml", "SRS parsing, developer matching, skill analysis endpoints"),
    ]
    for module, prefix, desc in routers:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(f"{module} ({prefix}): ")
        run.bold = True
        p.add_run(desc)
    add_heading(doc, "10.1 Middleware", 2)
    add_bullets(
        doc,
        [
            "CORSMiddleware: Allows frontend origins (localhost:5173, LAN IPs) with credentials.",
            "StripApiPrefixMiddleware: Accepts both /projects/... and /api/projects/... paths.",
        ],
    )

    # ── 11. ML/NLP ──
    add_heading(doc, "11. Machine Learning and NLP Components", 1)
    add_heading(doc, "11.1 Document Parser (document_parser.py)", 2)
    add_para(
        doc,
        "The DocumentParser class extracts text and requirements from uploaded SRS files. Supported "
        "formats: PDF (via pdfplumber with optional OCR), DOCX (via python-docx), TXT, and DOC (plain text fallback). "
        "It detects technical skills from a built-in dictionary, computes complexity scores, and "
        "suggests team sizes based on project complexity.",
    )
    add_heading(doc, "11.2 Skill Matcher (skill_matcher.py)", 2)
    add_para(
        doc,
        "SkillMatcher uses TF-IDF vectorization with character n-grams (3–5) and cosine similarity "
        "to compare required skills against developer skill profiles. The final match score combines "
        "weighted sub-scores:",
    )
    add_bullets(
        doc,
        [
            "Skill similarity: 40%",
            "Experience match: 20%",
            "Availability score: 15%",
            "CGPA score: 15%",
            "Performance history: 10%",
        ],
    )
    add_para(
        doc,
        "get_recommendations() ranks all available developers and returns the top matches with "
        "matching_skills lists and human-readable explanations via match_explanation.py.",
    )
    add_heading(doc, "11.3 Skill Extractor and Vectors", 2)
    add_para(
        doc,
        "ml/skill_extractor.py and ml/sklearn_vectors.py provide additional skill extraction and "
        "vector operations. ml/tensorflow_rank.py offers an optional TensorFlow-based ranker for "
        "advanced recommendation ordering.",
    )
    add_heading(doc, "11.4 Resume and Text Skill Extraction", 2)
    add_para(
        doc,
        "POST /ml/extract-skills-from-text accepts pasted text and returns detected skills. "
        "POST /users/{user_id}/resume accepts multipart resume uploads to merge extracted skills "
        "into the developer profile.",
    )

    # ── 12. Frontend ──
    add_heading(doc, "12. Frontend Application (React + Vite)", 1)
    add_heading(doc, "12.1 Routing Structure", 2)
    routes = [
        ("/login", "Login page (public only)"),
        ("/signup", "Registration page (public only)"),
        ("/dashboard", "Home dashboard (all authenticated users)"),
        ("/manager", "Manager/PM dashboard with project grid and AI recommendations"),
        ("/developer", "Developer tasks and skill profile page"),
        ("/analytics", "Analytics dashboard with charts and PDF download"),
        ("/developers", "Developers directory (manager/admin)"),
        ("/projects", "Projects directory"),
    ]
    for path, desc in routes:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(f"{path}: ")
        run.bold = True
        p.add_run(desc)
    add_heading(doc, "12.2 Key Components", 2)
    components = [
        ("Login.jsx / Signup.jsx", "Authentication forms with role-based redirect after login"),
        ("ManagerDashboard.jsx", "Active projects grid, completed projects section, team approval"),
        ("DeveloperDashboard.jsx", "Task list, skill profile, submission workflow"),
        ("DeveloperRecommendations.jsx", "AI recommendation cards with match explanations"),
        ("ProjectForm.jsx", "Project creation with optional SRS upload via FormData"),
        ("TaskList.jsx / ReviewSubmissions.jsx", "Task management and PM review queue"),
        ("AnalyticsDashboard.jsx", "Charts for utilization, gaps, team performance"),
        ("ChatBot.jsx", "Floating chatbot widget connected to POST /chatbot/query"),
        ("SkillProfile.jsx", "Developer skill proficiency display and editing"),
        ("ProtectedRoute.jsx", "Route guard enforcing authentication and role checks"),
        ("AuthContext.jsx", "Global auth state, token storage, user profile management"),
    ]
    for comp, desc in components:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(f"{comp}: ")
        run.bold = True
        p.add_run(desc)
    add_heading(doc, "12.3 API Integration", 2)
    add_para(
        doc,
        "services/api.js configures axios with base URL from VITE_API_URL (default /api for dev proxy). "
        "Request interceptors attach the JWT token. Response interceptors handle 401 redirects to login. "
        "Analytics requests use an extended timeout (VITE_ANALYTICS_TIMEOUT_MS, default 90 seconds).",
    )

    # ── 13. Workflows ──
    add_heading(doc, "13. Core Workflows", 1)
    add_heading(doc, "13.1 Project Creation Workflow", 2)
    add_numbered(
        doc,
        [
            "Manager logs in and navigates to the PM Dashboard.",
            "Creates a new project with title, description, deadline, department, and required skills.",
            "Optionally uploads an SRS document (PDF/DOCX/TXT) for automatic skill detection.",
            "Backend parses the document, merges detected skills, and stores the project in MongoDB.",
            "Project appears in the active projects grid.",
        ],
    )
    add_heading(doc, "13.2 AI Recommendation and Approval Workflow", 2)
    add_numbered(
        doc,
        [
            "Manager selects a project; frontend calls GET /projects/{id}/recommendations.",
            "SkillMatcher scores all available developers and returns top 5 with explanations.",
            "Manager reviews recommendation cards showing match scores and skill overlaps.",
            "Manager clicks Approve AI Team (POST /recommendations/{id}/approve).",
            "Backend updates assigned_team and final_team on the project document.",
            "Starter tasks are created for each team member, prioritizing developers with lower active workload.",
            "Developers see new task assignments on their dashboard.",
        ],
    )
    add_heading(doc, "13.3 Task Lifecycle Workflow", 2)
    add_numbered(
        doc,
        [
            "Task created with status assigned (by PM or automatically after team approval).",
            "Developer sets status to in_progress when starting work.",
            "Developer submits for review with optional file attachments.",
            "PM reviews submission in the Review Submissions queue.",
            "PM either requests changes (returns to in_progress) or marks as completed.",
            "On completion, developer skill proficiency is updated for skills_used on the task.",
        ],
    )
    add_heading(doc, "13.4 Adaptive Skill Update Workflow", 2)
    add_para(
        doc,
        "When a PM marks a task as completed, the tasks.py handler reads the skills_used array from "
        "the task document and bumps the assignee's skill proficiency in their user profile. This "
        "creates a feedback loop where completed work automatically enriches developer profiles for "
        "future AI matching.",
    )

    # ── 14. Recommendations ──
    add_heading(doc, "14. AI Team Recommendations", 1)
    add_para(
        doc,
        "The recommendation system is the core AI feature of the platform. It ensures fair, "
        "explainable, and manager-controlled team formation.",
    )
    add_bullets(
        doc,
        [
            "Stratified top-5 selection per project (MAX_PROJECT_TEAM_SIZE, default 5).",
            "Each candidate includes match_score, matching_skills, and explanation text.",
            "Approval creates recommendation records in the recommendations collection.",
            "No duplicate pending batches after approval (previous pending records are closed).",
            "Fair task distribution: sort_team_ids_for_fair_tasks orders by active task count.",
            "Team details modal shows resolved developer names, not raw MongoDB ObjectIds.",
            "Ghost/stale ObjectId rows are filtered from team display.",
        ],
    )

    # ── 15. Tasks ──
    add_heading(doc, "15. Task Management System", 1)
    add_bullets(
        doc,
        [
            "Tasks belong to a project and are assigned to a developer (assigned_to).",
            "Status values: assigned, in_progress, submitted, completed, and on_hold.",
            "Developers can submit work with file attachments stored in UPLOAD_DIR.",
            "PM review queue (GET /tasks/review-submissions) lists submitted tasks.",
            "Only managers/admins can set status to completed.",
            "Task descriptions can be auto-generated via task_description.py service.",
            "Maximum active tasks per developer enforced (default 10).",
            "Project metrics sync runs asynchronously after team approval.",
        ],
    )

    # ── 16. Analytics ──
    add_heading(doc, "16. Analytics and Reporting", 1)
    add_para(
        doc,
        "The analytics module (analytics.py and analytics_pdf.py) provides managers and admins with "
        "insights into team performance, skill utilization, and training needs.",
    )
    add_bullets(
        doc,
        [
            "Dashboard endpoint: GET /analytics/dashboard with rolling time windows (week/month/quarter/all).",
            "Utilization metrics showing developer workload distribution.",
            "Skill gap analysis comparing project requirements against team skills.",
            "Training recommendations for missing skills.",
            "Team performance summaries with completion rates.",
            "Per-project bar charts; falls back to all-time data if selected window has no activity.",
            "PDF performance report download (may take 30–90 seconds on large datasets).",
            "Frontend uses extended timeout (VITE_ANALYTICS_TIMEOUT_MS=90000) for heavy reports.",
        ],
    )

    # ── 17. Chatbot ──
    add_heading(doc, "17. Hybrid Chatbot", 1)
    add_heading(doc, "17.1 Architecture", 2)
    add_para(
        doc,
        "The chatbot (chat_engine.py + chatbot.py) uses a hybrid approach: a rule-based engine "
        "handles known intents first, and optionally falls back to a local Ollama LLM when the "
        "rule response would be too generic. MongoDB is the single source of truth — the LLM only "
        "reformulates data already loaded into JSON context.",
    )
    add_heading(doc, "17.2 Operating Modes", 2)
    add_bullets(
        doc,
        [
            "rules — Fast, deterministic answers from MongoDB only.",
            "hybrid — Rules first; Ollama fallback when answer is generic (recommended for demos).",
            "local_llm — Ollama first; rules fallback if model is unavailable.",
        ],
    )
    add_heading(doc, "17.3 Data Scope by Role", 2)
    add_bullets(
        doc,
        [
            "Admin: all projects, all tasks, all users.",
            "Manager/Developer: projects they own or are on; tasks on visible projects plus own assigned/created tasks.",
            "All users loaded from MongoDB (no hardcoded names); new users appear automatically.",
        ],
    )
    add_heading(doc, "17.4 Supported Query Types", 2)
    add_bullets(
        doc,
        [
            "Personal: my tasks, pending tasks, deadlines, today's plan.",
            "Project: status, summary, deadline, creator, task counts (single or all projects).",
            "Team: who is on project, team members, membership checks.",
            "Ownership: who handles module X, person Y's module work.",
            "Skills: my skills, proficiency checks, skill gap recommendations.",
            "Analytics: most active developer, highest completed tasks.",
            "Language: English and Roman Urdu/Hinglish with typo-tolerant name matching.",
        ],
    )
    add_heading(doc, "17.5 Privacy", 2)
    add_para(
        doc,
        "By default, Ollama runs locally (http://127.0.0.1:11434) — no cloud LLM required. "
        "This ensures project data stays on-premises during demos and production use.",
    )
    add_heading(doc, "17.6 Evaluation Battery", 2)
    add_para(
        doc,
        "A 400-query test battery (200 English + 200 Roman Urdu) is available via "
        "scripts/run_chatbot_battery.py for systematic chatbot evaluation and jury testing.",
    )

    # ── 18. API Overview ──
    add_heading(doc, "18. API Endpoints Overview", 1)
    add_para(doc, "Key API endpoints grouped by module (full documentation at http://127.0.0.1:8000/docs):")
    endpoints = [
        ("Authentication", "POST /auth/register, POST /auth/login, GET /auth/me"),
        ("Users", "GET /users/developers, PUT /users/{id}, POST /users/{id}/skills, POST /users/{id}/resume"),
        ("Projects", "POST /projects/, GET /projects/, GET /projects/{id}, PUT /projects/{id}, GET /projects/{id}/recommendations"),
        ("Tasks", "GET /tasks/, POST /tasks/, PUT /tasks/{id}/status, POST /tasks/{id}/submit-for-review, GET /tasks/review-submissions"),
        ("Recommendations", "POST /recommendations/, POST /recommendations/{id}/approve, POST /recommendations/{id}/reject"),
        ("ML", "POST /ml/parse-srs, POST /ml/match-developers, POST /ml/extract-skills-from-text, GET /ml/skill-analysis"),
        ("Analytics", "GET /analytics/dashboard, GET /analytics/performance-report"),
        ("Chatbot", "POST /chatbot/query"),
        ("Teams", "GET /teams/, team management endpoints"),
    ]
    for group, eps in endpoints:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(f"{group}: ")
        run.bold = True
        p.add_run(eps)

    # ── 19. Environment ──
    add_heading(doc, "19. Environment Configuration", 1)
    add_para(doc, "Backend environment variables (backend/.env, template in .env.example):")
    env_vars = [
        ("MONGO_URI / MONGODB_URL", "MongoDB connection string"),
        ("MONGODB_DB_NAME", "Database name (default: skill_mapping)"),
        ("SECRET_KEY", "JWT signing key (must change in production)"),
        ("ACCESS_TOKEN_EXPIRE_MINUTES", "Token lifetime (default: 10080 = 7 days)"),
        ("MAX_PROJECT_TEAM_SIZE", "AI team roster size (default: 5)"),
        ("CHATBOT_MODE", "rules | hybrid | local_llm"),
        ("OLLAMA_BASE_URL", "Ollama server URL (default: http://127.0.0.1:11434)"),
        ("OLLAMA_MODEL", "Model name (e.g., gemma3:27b)"),
        ("OLLAMA_TIMEOUT_SECONDS", "LLM request timeout"),
        ("UPLOAD_DIR", "File upload directory"),
        ("MAX_FILE_SIZE", "Maximum upload size in bytes"),
    ]
    for var, desc in env_vars:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(f"{var}: ")
        run.bold = True
        p.add_run(desc)
    add_para(doc, "Frontend environment variables (frontend/.env):")
    add_bullets(
        doc,
        [
            "VITE_API_URL=/api (uses Vite dev proxy to backend)",
            "VITE_API_TIMEOUT_MS=25000 (general API timeout)",
            "VITE_ANALYTICS_TIMEOUT_MS=90000 (analytics/PDF timeout)",
        ],
    )

    # ── 20. Installation ──
    add_heading(doc, "20. Installation and Running the Platform", 1)
    add_heading(doc, "20.1 Prerequisites", 2)
    add_bullets(
        doc,
        [
            "Python 3.11+ with virtual environment",
            "Node.js 18+",
            "MongoDB (local, Docker, or Atlas)",
            "Optional: Ollama for hybrid chatbot",
        ],
    )
    add_heading(doc, "20.2 Backend Setup", 2)
    add_para(doc, "From the backend directory:")
    code_steps = [
        "Copy .env.example to .env and configure MongoDB URI and SECRET_KEY.",
        "Create virtual environment: python -m venv venv312",
        "Activate: .\\venv312\\Scripts\\activate (Windows)",
        "Install dependencies: pip install -r requirements.txt",
        "Seed database (optional): python seed_database.py",
        "Start server: python -m uvicorn app.main:app --host 127.0.0.1 --port 8000 --reload",
        "Verify: open http://127.0.0.1:8000/docs",
    ]
    add_numbered(doc, code_steps)
    add_heading(doc, "20.3 Frontend Setup", 2)
    add_para(doc, "From the frontend directory:")
    fe_steps = [
        "npm install",
        "npm run dev",
        "Open the URL printed by Vite (typically http://localhost:5173)",
    ]
    add_numbered(doc, fe_steps)
    add_heading(doc, "20.4 Demo Accounts", 2)
    add_para(doc, "After seeding, example accounts include:")
    add_bullets(
        doc,
        [
            "Manager: manager@demo.local / manager123",
            "Developer: arooj123@gmail.com / arooj123 (or aima123@gmail.com, haseeb123@gmail.com)",
            "Admin (CEO): mrehaansaleemceo123@gmail.com / Pakistan123",
        ],
    )

    # ── 21. Demo Scenarios ──
    add_heading(doc, "21. Demo Scenarios", 1)
    scenarios = [
        ("Login and AI Recommendations", "Log in as manager, select project, view top-5 developer recommendations with explanations."),
        ("Approve AI Team", "Approve recommendations; verify MongoDB team fields and starter task creation."),
        ("ML Developer Matching", "Use Swagger POST /ml/match-developers with required skills array."),
        ("SRS Skill Extraction", "Upload PDF/DOCX via POST /ml/parse-srs or paste text for skill extraction."),
        ("Adaptive Skills", "Complete a task as PM; verify developer skill proficiency increases."),
        ("Chatbot Demo", "Ask questions in English or Roman Urdu; show live MongoDB grounding."),
        ("Analytics Dashboard", "View charts and download PDF performance report."),
        ("RBAC Smoke Test", "Verify developers get 403 on restricted endpoints."),
    ]
    for i, (name, desc) in enumerate(scenarios, 1):
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(f"{name}: ")
        run.bold = True
        p.add_run(desc)

    # ── 22. Security ──
    add_heading(doc, "22. Security and Privacy Considerations", 1)
    add_bullets(
        doc,
        [
            "Passwords hashed with bcrypt; never stored in plain text.",
            "JWT tokens with configurable expiration; SECRET_KEY must be changed for production.",
            "Role-based access control on all sensitive endpoints.",
            "CORS restricted to configured origins and LAN regex pattern.",
            "Chatbot uses local Ollama by default — no external API calls for LLM.",
            "Password fields stripped from chatbot user context JSON.",
            "File uploads validated by extension and size limits.",
            "Demo project titles hidden from production UI to avoid confusion.",
        ],
    )

    # ── 23. Limitations ──
    add_heading(doc, "23. Limitations and Future Enhancements", 1)
    add_heading(doc, "23.1 Current Limitations", 2)
    add_bullets(
        doc,
        [
            "Skill matching relies on string-based TF-IDF; semantic embeddings could improve accuracy.",
            "Ollama LLM requires local GPU/CPU resources for larger models.",
            "Analytics PDF generation can be slow on large MongoDB datasets.",
            "Single-organization design; no multi-tenant support yet.",
        ],
    )
    add_heading(doc, "23.2 Future Enhancements", 2)
    add_bullets(
        doc,
        [
            "Integration with sentence-transformers for semantic skill matching.",
            "Real-time notifications via WebSockets.",
            "Mobile-responsive PWA or native mobile app.",
            "Integration with GitHub/GitLab for automatic skill inference from commits.",
            "Multi-language UI beyond chatbot (full Urdu interface).",
            "Automated sprint planning based on skill gaps and deadlines.",
        ],
    )

    # ── 24. Conclusion ──
    add_heading(doc, "24. Conclusion", 1)
    add_para(
        doc,
        "The Skill Mapping Platform delivers a complete SDS-compliant loop for AI-powered developer "
        "team formation and skill management. From document parsing and ML-based matching through "
        "manager approval, task tracking, adaptive profiling, analytics, and a bilingual hybrid chatbot, "
        "the system addresses the full lifecycle of skill-aware project management.",
    )
    add_para(
        doc,
        "Built with modern technologies (FastAPI, React, MongoDB, scikit-learn, Ollama), the platform "
        "is designed for academic demonstration, viva defense, and practical deployment in "
        "software development organizations. Group F25CS093 has implemented a production-quality "
        "architecture with explainable AI, role-based security, and privacy-preserving local LLM "
        "integration — making it a comprehensive Final Year Project solution.",
    )

    return doc


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.save(str(OUT_FILE))
    print(f"Saved: {OUT_FILE}")


if __name__ == "__main__":
    main()
